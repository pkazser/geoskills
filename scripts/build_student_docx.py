"""Build GeoSkills student DOCX files from Markdown source files.

The Markdown files remain the single content source. This lightweight converter
supports the structures used by the GeoSkills student sheets: headings, bold
labels, blockquotes, tables and answer areas.
"""

from __future__ import annotations

import argparse
import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


BLUE = "1F4E79"
LIGHT_BLUE = "D9EAF7"
GRID = "B7C9D6"
GRAY = "666666"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_borders(cell, color: str = GRID) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_padding(cell, value: int = 100) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge in ("top", "start", "bottom", "end"):
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("Σελίδα ")
    run.font.size = Pt(9)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    paragraph._p.append(field)


def set_run_font(run, size: float | None = None, bold: bool | None = None) -> None:
    run.font.name = "Arial"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Arial")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Arial")
    run._element.get_or_add_rPr().rFonts.set(qn("w:cs"), "Arial")
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def add_inline(paragraph, text: str, size: float = 10.5) -> None:
    parts = re.split(r"(\*\*.*?\*\*)", text)
    for part in parts:
        if not part:
            continue
        bold = part.startswith("**") and part.endswith("**")
        value = part[2:-2] if bold else part
        run = paragraph.add_run(value)
        set_run_font(run, size=size, bold=bold)


def is_table_line(line: str) -> bool:
    return line.strip().startswith("|") and line.strip().endswith("|")


def table_rows(lines: list[str]) -> list[list[str]]:
    rows = []
    for line in lines:
        cells = [cell.strip() for cell in line.strip().strip("|").split("|")]
        if all(re.fullmatch(r":?-{3,}:?", cell) for cell in cells):
            continue
        rows.append(cells)
    return rows


def add_table(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    cols = max(len(row) for row in rows)
    table = doc.add_table(rows=len(rows), cols=cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for r, row in enumerate(rows):
        for c in range(cols):
            cell = table.cell(r, c)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_borders(cell)
            set_cell_padding(cell)
            if c < len(row):
                cell.text = ""
                p = cell.paragraphs[0]
                p.paragraph_format.space_after = Pt(2)
                add_inline(p, row[c], size=9.3)
            if r == 0:
                set_cell_shading(cell, BLUE)
                for run in cell.paragraphs[0].runs:
                    run.font.color.rgb = RGBColor(255, 255, 255)
                    run.bold = True
            elif r % 2 == 0:
                set_cell_shading(cell, "F5F9FC")
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_answer_space(doc: Document, lines: int = 3) -> None:
    """Actual editable ruled rows; consecutive paragraph borders can collapse on render."""
    table = doc.add_table(rows=lines, cols=1)
    table.autofit = False
    for row in table.rows:
        row.height = Cm(0.51)
        row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
        cell = row.cells[0]
        cell.width = Cm(17)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_padding(cell, value=10)
        tc_pr = cell._tc.get_or_add_tcPr()
        borders = tc_pr.first_child_found_in("w:tcBorders")
        if borders is None:
            borders = OxmlElement("w:tcBorders")
            tc_pr.append(borders)
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "4")
        bottom.set(qn("w:color"), GRID)
        borders.append(bottom)
        paragraph = cell.paragraphs[0]
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 1.0
        run = paragraph.add_run("\u00a0")
        set_run_font(run, size=8.5)


def answer_lines(document_type: str, line: str, question_number: int | None) -> int:
    """Return a deliberate answer-space size for the specific worksheet."""
    if document_type == "06":
        return 0
    if document_type == "05":
        if "Τεκμηρίωση απόφασης" in line:
            return 4
        if "Επόμενος έλεγχος" in line:
            return 2
        # Existing dotted placeholders are already editable fields; do not double them.
        return 0
        return 0
    if document_type == "07" and question_number is not None and line.strip().startswith("Απάντηση"):
        sizes = {1: 3, 2: 4, 4: 4, 5: 3, 6: 5, 7: 4, 8: 4, 9: 4, 10: 4, 11: 4, 12: 4, 13: 5}
        return sizes.get(question_number, 3)
    return 0


def apply_styles(doc: Document) -> None:
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.08
    for name, size, color in (("Title", 19, "000000"), ("Heading 1", 14, "000000"), ("Heading 2", 12, BLUE), ("Heading 3", 11, BLUE)):
        style = styles[name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(10 if name != "Title" else 0)
        style.paragraph_format.space_after = Pt(5)


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.7)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)
    footer = section.footer
    add_page_number(footer.paragraphs[0])


def build(source: Path, output: Path, code: str, version: str) -> None:
    lines = source.read_text(encoding="utf-8").splitlines()
    doc = Document()
    apply_styles(doc)
    document_type = code.rsplit(" ", 1)[-1]
    configure_document(doc)
    if document_type == "05":
        doc.sections[0].top_margin = Cm(1.5)
        doc.sections[0].bottom_margin = Cm(1.4)
    if document_type == "07":
        doc.styles["Normal"].font.size = Pt(9.7)
        doc.styles["Normal"].paragraph_format.space_after = Pt(3)
        doc.sections[0].top_margin = Cm(1.45)
        doc.sections[0].bottom_margin = Cm(1.35)
    table_buffer: list[str] = []
    title_seen = False
    question_number: int | None = None
    skip_single_dotted_placeholder = False

    def flush_table() -> None:
        nonlocal table_buffer
        if table_buffer:
            add_table(doc, table_rows(table_buffer))
            table_buffer = []

    for raw in lines:
        line = raw.rstrip()
        if is_table_line(line):
            table_buffer.append(line)
            continue
        flush_table()
        if not line.strip():
            continue
        if skip_single_dotted_placeholder:
            skip_single_dotted_placeholder = False
            if re.fullmatch(r"\.+", line.strip()):
                continue
        heading = re.match(r"^(#{1,3})\s+(.*)$", line)
        if heading:
            level = len(heading.group(1))
            text = heading.group(2).strip()
            p = doc.add_paragraph(style="Title" if not title_seen and level == 1 else f"Heading {level}")
            p.paragraph_format.keep_with_next = True
            if document_type == "07" and text.startswith("Ε. Ερμηνεία"):
                p.paragraph_format.page_break_before = True
            add_inline(p, text, size=19 if not title_seen and level == 1 else 11 + (3 - level))
            title_seen = True
            continue
        if line.startswith(">"):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.5)
            p.paragraph_format.space_after = Pt(8)
            run = p.add_run(line[1:].strip())
            set_run_font(run, size=10, bold=False)
            run.italic = True
            run.font.color.rgb = RGBColor.from_string(GRAY)
            continue
        if line.startswith("- [ ]"):
            p = doc.add_paragraph(style="Normal")
            p.paragraph_format.left_indent = Cm(0.35)
            add_inline(p, "☐ " + line[5:].strip())
            continue
        if re.match(r"^\*\*.*\*\*$", line):
            p = doc.add_paragraph()
            p.paragraph_format.keep_with_next = True
            add_inline(p, line)
            match = re.match(r"^\*\*(\d+)\.", line)
            if match:
                question_number = int(match.group(1))
                if document_type == "07" and question_number == 6:
                    p.paragraph_format.page_break_before = True
            if document_type == "05" and ("Τεκμηρίωση απόφασης" in line or "Επόμενος έλεγχος" in line):
                count = answer_lines(document_type, line, None)
                add_answer_space(doc, count)
                skip_single_dotted_placeholder = True
            continue
        p = doc.add_paragraph()
        if document_type == "07" and ("Απάντηση" in line or "Απάντηση /" in line):
            p.paragraph_format.keep_with_next = True
        add_inline(p, line)
        count = answer_lines(document_type, line, question_number)
        if count:
            add_answer_space(doc, count)

    flush_table()
    first = doc.paragraphs[0]
    meta = doc.paragraphs[1].insert_paragraph_before(f"Κωδικός ενότητας: {code}  |  Έκδοση: {version}")
    meta.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for run in meta.runs:
        set_run_font(run, size=9.5, bold=True)
        run.font.color.rgb = RGBColor.from_string(BLUE)
    doc.core_properties.title = f"GeoSkills {code} student document"
    doc.core_properties.subject = "GeoSkills student material"
    doc.core_properties.author = "GeoSkills"
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--code", required=True)
    parser.add_argument("--version", default="1.0")
    args = parser.parse_args()
    build(args.source, args.output, args.code, args.version)


if __name__ == "__main__":
    main()
